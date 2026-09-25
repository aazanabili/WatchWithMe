from pathlib import Path
import subprocess

from docx import Document
from docx.shared import Mm
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
SOURCE = Path(r"C:\Users\dellaaz\Downloads\Documents\Lebenslauf Mohanad Alkawarid.pdf")
POPPLER = Path(r"C:\Users\dellaaz\.cache\codex-runtimes\codex-primary-runtime\dependencies\native\poppler\Library\bin\pdftoppm.exe")
RENDER_DIR = ROOT / "tmp" / "cv_600dpi"
OUTPUT = ROOT / "Lebenslauf Mohanad Alkawarid.docx"

RENDER_DIR.mkdir(parents=True, exist_ok=True)
subprocess.run(
    [str(POPPLER), "-png", "-r", "600", str(SOURCE), str(RENDER_DIR / "page")],
    check=True,
)

doc = Document()
doc.add_paragraph()
section = doc.sections[0]
section.page_width = Mm(210)
section.page_height = Mm(297)
section.top_margin = Mm(0)
section.bottom_margin = Mm(0)
section.left_margin = Mm(0)
section.right_margin = Mm(0)
section.header_distance = Mm(0)
section.footer_distance = Mm(0)

for index, image in enumerate(sorted(RENDER_DIR.glob("page-*.png"))):
    if index:
        section = doc.add_section(WD_SECTION_START.NEW_PAGE)
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(0)
        section.bottom_margin = Mm(0)
        section.left_margin = Mm(0)
        section.right_margin = Mm(0)
        section.header_distance = Mm(0)
        section.footer_distance = Mm(0)

    paragraph = doc.paragraphs[-1] if index == 0 else doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Mm(0)
    paragraph.paragraph_format.space_after = Mm(0)
    paragraph.paragraph_format.line_spacing = 1
    run = paragraph.add_run()
    run.add_picture(str(image), width=Mm(210), height=Mm(297))
    ppr = paragraph._p.get_or_add_pPr()
    spacing = ppr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        ppr.append(spacing)
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), "240")
    spacing.set(qn("w:lineRule"), "auto")

doc.core_properties.title = "Lebenslauf Mohanad Alkawarid"
doc.save(OUTPUT)
print(OUTPUT)
